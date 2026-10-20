# SPDX-License-Identifier: MIT
# Copyright (c) 2024-2026 RAEN Digital Tools SL - PyNET Platform

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from System import Environment

# The host injects the Presto COM application (Presto.App.18) as __presto__
try:
    app = __presto__
except NameError:
    raise Exception("__presto__ not found: run this script from PyNET for Presto (Complementos > PyNET).")

# ---------------------------------------------------------------------------
# Export the budget tree of the open obra to Excel: one row per concept with
# level, parent code, code, nature, unit, summary, quantity, unit price and
# amount, in an Excel table, saved to the Desktop as "<obra> - Budget.xlsx".
#
# Same walk as WalkBudgetTree.py (one element slot per depth, _SPC_ and _CHG_
# filtered out, only chapters are expanded unless DESCEND_INTO_DECOMPOSITION).
# Excel rules from RevitApiPatterns: no column auto-fit via ws.columns,
# table named "Tabla_1" (never a cell-like name such as "T1"), empty cells
# written as None.
#
# Read-only on Presto. Verified on Presto 26.01 (2026-10-20, whole script):
# tree walk, codes, summaries, Nat, CanPres, Conceptos.Ud, Conceptos.Pres
# (unit price) and Relaciones.ImpPres (amount); table "Tabla_1" opens in Excel.
# ---------------------------------------------------------------------------
OUTPUT_DIR = Path(Environment.GetFolderPath(Environment.SpecialFolder.Desktop))
MAX_DEPTH = 8
DESCEND_INTO_DECOMPOSITION = False
FILTER = "Conceptos.Nat!=219 && Conceptos.Nat!=170"

NAT_ROOT = 0
NAT_CHAPTER = 4
NATURES = {0: "Root", 4: "Chapter", 5: "Partida", 6: "Labour", 7: "Material", 8: "Machinery", 9: "Other"}
HEADERS = ["Level", "Parent", "Code", "Nature", "Unit", "Summary", "Quantity", "Price", "Amount"]


def num(field):
    value = app.GetFieldNum(field)          # numeric fields: GetFieldNum, never GetFieldStr
    return float(value) if value is not None else None


def walk(sup, depth, rows):
    app.SetElement(depth, "Relaciones", "Relaciones.CodSup", '"' + sup + '"', FILTER)
    while app.GetElement(depth) == 0:
        code = app.GetFieldStr("Relaciones.CodInf")
        nat = int(num("Conceptos.Nat") or 0)
        rows.append([
            depth, sup, code, NATURES.get(nat, str(nat)),
            app.GetFieldStr("Conceptos.Ud") or None,
            app.GetFieldStr("Conceptos.Resumen") or None,
            num("Relaciones.CanPres"), num("Conceptos.Pres"), num("Relaciones.ImpPres"),
        ])
        if depth < MAX_DEPTH and (nat == NAT_CHAPTER or DESCEND_INTO_DECOMPOSITION):
            walk(code, depth + 1, rows)


project = str(app.Project() or "")
if not project:
    raise Exception("No obra is open in Presto: open one and run the script again.")
if app.FindEqual("Conceptos", "Conceptos.Nat", NAT_ROOT) != 0:
    raise Exception("The obra has no root concept (Nat 0).")
root = app.GetFieldStr("Conceptos.Código")
root_summary = app.GetFieldStr("Conceptos.Resumen") or None
budget = app.EvalNum("Conceptos.Pres")

rows = []
walk(root, 1, rows)
if not rows:
    raise Exception("The budget is empty: nothing to export.")

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(HEADERS)
ws.append([0, None, root, NATURES[NAT_ROOT], None, root_summary, 1, budget, budget])
for row in rows:
    ws.append([v if v != "" else None for v in row])

bold = Font(bold=True)
for row_cells in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=len(HEADERS)):
    if row_cells[3].value == "Chapter":
        for cell in row_cells:
            cell.font = bold
for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=7, max_col=9):
    for cell in row_cells:
        cell.number_format = "#,##0.00"

table = Table(displayName="Tabla_1", ref=f"A1:I{ws.max_row}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
ws.freeze_panes = "A2"

out = OUTPUT_DIR / f"{Path(project).stem} - Budget.xlsx"
try:
    wb.save(out)
except PermissionError:
    raise Exception(f"Cannot write {out}: close it in Excel and run the script again.")

print(f"Exported {len(rows)} concepts to {out}")
ia_Result = {"type": "PrestoExport", "file": str(out), "rows": len(rows), "budget": budget}
