# SPDX-License-Identifier: MIT
# Copyright (c) 2024-2026 RAEN Digital Tools SL - PyNET Platform

import csv
from collections import Counter, OrderedDict
from pathlib import Path

from System import Environment

# The host injects the Presto COM application (Presto.App.18) as __presto__
try:
    app = __presto__
except NameError:
    raise Exception("__presto__ not found: run this script from PyNET for Presto (Complementos > PyNET).")

# ---------------------------------------------------------------------------
# ISO 19650 traceability check of the measurement lines (table Mediciones)
# of the open obra: for every partida (tree walk like WalkBudgetTree.py,
# _SPC_/_CHG_ skipped) or only the CODES listed:
#   - number of lines, sum of Mediciones.Cantidad (lines with Pres = 1)
#     against the partida quantity Relaciones.CanPres -> mismatches flagged;
#   - lines without GuidElem (no link to a model element);
#   - GuidElem repeated across the obra (an element measured twice);
#   - lines per Planta and per CategoríaBIM.
# Optional export of every line to the Desktop (EXPORT = "csv" or "xlsx").
# Read-only, changes nothing.
#
# Lines written by WriteMeasurementLinesFromJson.py (from a Tekla / BIM
# takeoff, one line per element with its IFC GlobalId) should give: no line
# without GuidElem, no duplicates, sum == CanPres for every partida.
#
# Verified on Presto 26.01 (2026-10-20), the pieces: lines walked with
# SetElement(slot, "Mediciones", "Mediciones.Relación", '"*"',
# 'Mediciones.CodSup=="X" && Mediciones.CodInf=="Y"') + GetElement; parents
# of a partida with SetElement on Relaciones and mask '"*"'; Cantidad /
# CanPres read with GetFieldNum; text fields with GetFieldStr.
# Whole script run on Presto 26.01 (2026-10-20, EXPORT "xlsx"): 570 lines /
# 10 partidas in ~15 s, all checks and the Excel export OK.
# ---------------------------------------------------------------------------
CODES = []                     # [] -> every partida of the tree; e.g. ["01.01", "02.05"]
EXPORT = ""                    # "", "csv" or "xlsx" (Desktop, "<obra> - Mediciones.<ext>")
TOLERANCE = 0.001              # |sum of lines - CanPres| allowed
MAX_DEPTH = 8                  # slots 1..MAX_DEPTH for the tree walk
LINE_SLOT = 9                  # slot for the line / relation lookups (after the walk)
FILTER = "Conceptos.Nat!=219 && Conceptos.Nat!=170"
NAT_ROOT, NAT_CHAPTER, NAT_PARTIDA = 0, 4, 5

TEXT_FIELDS = ["GuidElem", "Comentario", "Planta", "CategoríaBIM", "TipoBIM", "FamiliaBIM", "Área", "Archivo", "Nota"]
NUM_FIELDS = ["N", "Longitud", "Anchura", "Altura", "Cantidad", "Pres", "Cota", "BIMLong", "BIMVol"]


def q(text):
    return '"' + text + '"'


def num(field):
    value = app.GetFieldNum(field)          # numeric fields: GetFieldStr returns None
    return float(value) if value is not None else None


def walk(sup, depth, out):
    app.SetElement(depth, "Relaciones", "Relaciones.CodSup", q(sup), FILTER)
    while app.GetElement(depth) == 0:
        code = app.GetFieldStr("Relaciones.CodInf")
        nat = int(num("Conceptos.Nat") or 0)
        if nat == NAT_PARTIDA:
            out.append({"sup": sup, "code": code, "summary": app.GetFieldStr("Conceptos.Resumen") or "",
                        "unit": app.GetFieldStr("Conceptos.Ud") or "", "canpres": num("Relaciones.CanPres") or 0.0})
        elif nat == NAT_CHAPTER and depth < MAX_DEPTH:
            walk(code, depth + 1, out)


def relations_of(code):
    """Every parent -> code relation (a partida may hang from several parents)."""
    out = []
    app.SetElement(LINE_SLOT, "Relaciones", "Relaciones.CodSup", '"*"', "Relaciones.CodInf==" + q(code))
    while app.GetElement(LINE_SLOT) == 0:
        out.append({"sup": app.GetFieldStr("Relaciones.CodSup"), "code": code,
                    "summary": app.GetFieldStr("Conceptos.Resumen") or "",
                    "unit": app.GetFieldStr("Conceptos.Ud") or "", "canpres": num("Relaciones.CanPres") or 0.0})
    return out


def read_lines(sup, code):
    rows = []
    app.SetElement(LINE_SLOT, "Mediciones", "Mediciones.Relación", '"*"',
                   "Mediciones.CodSup==" + q(sup) + " && Mediciones.CodInf==" + q(code))
    while app.GetElement(LINE_SLOT) == 0:
        row = {"CodSup": sup, "CodInf": code}
        for f in TEXT_FIELDS:
            row[f] = app.GetFieldStr("Mediciones." + f) or ""
        for f in NUM_FIELDS:
            row[f] = num("Mediciones." + f)
        rows.append(row)
    return rows


project = str(app.Project() or "")
if not project:
    raise Exception("No obra is open in Presto: open one and run the script again.")

relations = []
if CODES:
    for code in CODES:
        found = relations_of(code)
        if not found:
            print(f"{code}: not found in the budget tree")
        relations.extend(found)
else:
    if app.FindEqual("Conceptos", "Conceptos.Nat", NAT_ROOT) != 0:
        raise Exception("The obra has no root concept (Nat 0).")
    walk(app.GetFieldStr("Conceptos.Código"), 1, relations)

print(f"{project}: reading the lines of {len(relations)} partidas...")
all_lines, report = [], []
guid_where = {}
for rel in relations:
    lines = read_lines(rel["sup"], rel["code"])
    counted = [l for l in lines if (l["Pres"] if l["Pres"] is not None else 1) == 1]
    total = sum(l["Cantidad"] or 0.0 for l in counted)
    no_guid = sum(1 for l in lines if not l["GuidElem"].strip())
    for l in lines:
        g = l["GuidElem"].strip()
        if g:
            guid_where.setdefault(g, []).append(rel["sup"] + "/" + rel["code"])
    mismatch = bool(lines) and abs(total - rel["canpres"]) > TOLERANCE
    report.append({"type": "PrestoPartidaLines", "parent": rel["sup"], "code": rel["code"], "unit": rel["unit"],
                   "summary": rel["summary"], "lines": len(lines), "lines_not_counted": len(lines) - len(counted),
                   "sum_lines": round(total, 4), "canpres": rel["canpres"], "mismatch": mismatch,
                   "lines_without_guid": no_guid})
    all_lines.extend(lines)

duplicates = OrderedDict((g, w) for g, w in guid_where.items() if len(w) > 1)
by_storey = Counter(l["Planta"] or "(none)" for l in all_lines)
by_class = Counter(l["CategoríaBIM"] or "(none)" for l in all_lines)
with_lines = [r for r in report if r["lines"]]

for r in with_lines:
    flag = "  <-- MISMATCH" if r["mismatch"] else ""
    flag += f"  ({r['lines_without_guid']} without GuidElem)" if r["lines_without_guid"] else ""
    print(f"  {r['parent']}/{r['code']}  {r['lines']:>5} lines  sum {r['sum_lines']:,.3f}  "
          f"CanPres {r['canpres']:,.3f} {r['unit']}{flag}")
print(f"Partidas: {len(report)} | with lines: {len(with_lines)} | lines: {len(all_lines)} | "
      f"without GuidElem: {sum(r['lines_without_guid'] for r in report)} | duplicated GuidElem: {len(duplicates)} | "
      f"sum mismatches: {sum(1 for r in report if r['mismatch'])}")
print("By Planta: " + ", ".join(f"{k}: {v}" for k, v in sorted(by_storey.items())))
print("By CategoríaBIM: " + ", ".join(f"{k}: {v}" for k, v in by_class.most_common()))

exported = None
if EXPORT and all_lines:
    headers = ["CodSup", "CodInf"] + TEXT_FIELDS + NUM_FIELDS
    out = Path(Environment.GetFolderPath(Environment.SpecialFolder.Desktop)) / \
        f"{Path(project).stem} - Mediciones.{EXPORT.lower()}"
    try:
        if EXPORT.lower() == "csv":
            with open(out, "w", newline="", encoding="utf-8-sig") as fh:     # BOM: Excel reads the accents
                writer = csv.writer(fh, delimiter=";")
                writer.writerow(headers)
                for l in all_lines:
                    writer.writerow([l[h] for h in headers])
        else:
            from openpyxl import Workbook
            from openpyxl.worksheet.table import Table, TableStyleInfo
            wb = Workbook()
            ws = wb.active
            ws.title = "Mediciones"
            ws.append(headers)
            for l in all_lines:
                ws.append([l[h] if l[h] != "" else None for h in headers])      # empty cells as None
            table = Table(displayName="Tabla_1", ref=f"A1:{chr(64 + len(headers))}{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
            ws.freeze_panes = "A2"
            wb.save(out)
    except PermissionError:
        raise Exception(f"Cannot write {out}: close it and run the script again.")
    exported = str(out)
    print(f"Exported {len(all_lines)} lines to {out}")

ia_Result = {"type": "PrestoMeasurementLinesCheck", "project": project, "partidas": report,
             "lines": len(all_lines), "lines_without_guid": sum(r["lines_without_guid"] for r in report),
             "duplicated_guids": [{"guid": g, "in": w} for g, w in duplicates.items()],
             "mismatches": [r["code"] for r in report if r["mismatch"]],
             "by_storey": dict(by_storey), "by_ifc_class": dict(by_class), "file": exported}
