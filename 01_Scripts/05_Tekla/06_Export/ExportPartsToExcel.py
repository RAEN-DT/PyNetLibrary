# SPDX-License-Identifier: MIT
# Copyright (c) 2024-2026 RAEN Digital Tools SL - PyNET Platform

import clr
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

clr.AddReference("Tekla.Structures.Model")

from System import Array, Environment, Int32, Type
from Tekla.Structures.Model import Model, Beam, ContourPlate, PolyBeam

# The host injects the connected model as __teklamodel__
try:
    model = __teklamodel__
except NameError:
    model = Model()

# ---------------------------------------------------------------------------
# Export the part inventory to Excel: one row per Beam / column, ContourPlate
# and PolyBeam with GUID, type, name, class, profile, material, assembly and
# part position, phase, length (mm) and weight (kg), in an Excel table saved
# to the Desktop as "<model> - Parts.xlsx". Read-only on the model.
#
# Same read as ListParts.py. Excel rules (RevitApiPatterns "Exporting to
# Excel"): no column auto-fit via ws.columns, table named "Tabla_1" (never a
# cell-like name such as "T1"), empty cells written as None. The Desktop path
# comes from Environment.GetFolderPath (OneDrive may redirect it).
#
# Verified on Tekla 2026 (2026-10-16): the whole script - GetAllObjectsWithType,
# GetReportProperty (found, value) tuples for LENGTH / WEIGHT / ASSEMBLY_POS /
# PART_POS / PHASE (Int32 default), GetSize(), the openpyxl table.
# ---------------------------------------------------------------------------
OUTPUT_DIR = Path(Environment.GetFolderPath(Environment.SpecialFolder.Desktop))
PART_TYPES = [Beam, ContourPlate, PolyBeam]
HEADERS = ["GUID", "Type", "Name", "Class", "Profile", "Material",
           "Assembly pos", "Part pos", "Phase", "Length (mm)", "Weight (kg)"]


def report(obj, name, default):
    """Report property value, or None when Tekla has none for this object."""
    found, value = obj.GetReportProperty(name, default)
    return value if found else None


if not model.GetConnectionStatus():
    raise Exception("Not connected to Tekla Structures.")

types = Array[Type]([clr.GetClrType(t) for t in PART_TYPES])
enumerator = model.GetModelObjectSelector().GetAllObjectsWithType(types)
total = enumerator.GetSize()
step = max(1, total // 10)
print(f"Reading {total} parts...")

rows = []
while enumerator.MoveNext():
    part = enumerator.Current
    if part is None:
        continue
    if len(rows) and len(rows) % step == 0:
        print(f"  {len(rows)}/{total} ({100 * len(rows) // max(total, 1)}%)")

    kind = part.GetType().Name
    if isinstance(part, Beam):
        kind = "Column" if str(part.Type) == "COLUMN" else "Beam"
    length = report(part, "LENGTH", 0.0)
    weight = report(part, "WEIGHT", 0.0)
    rows.append([
        str(part.Identifier.GUID),
        kind,
        part.Name or None,
        part.Class or None,
        part.Profile.ProfileString or None,
        part.Material.MaterialString or None,
        report(part, "ASSEMBLY_POS", "") or None,
        report(part, "PART_POS", "") or None,
        report(part, "PHASE", Int32(0)),   # integer property: a Python 0 picks the Double overload
        round(length, 1) if length is not None else None,
        round(weight, 2) if weight is not None else None,
    ])

if not rows:
    raise Exception("The model has no parts to export.")

wb = Workbook()
ws = wb.active
ws.title = "Parts"
ws.append(HEADERS)
for row in rows:
    ws.append(row)
for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=10, max_col=11):
    for cell in row_cells:
        cell.number_format = "#,##0.0"

table = Table(displayName="Tabla_1", ref=f"A1:K{ws.max_row}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
ws.freeze_panes = "A2"

model_name = Path(model.GetInfo().ModelName).stem      # "X.db1" -> "X"
out = OUTPUT_DIR / f"{model_name} - Parts.xlsx"
try:
    wb.save(out)
except PermissionError:
    raise Exception(f"Cannot write {out}: close it in Excel and run the script again.")

print(f"Exported {len(rows)} parts to {out}")
ia_Result = {"type": "TeklaPartsExport", "file": str(out), "rows": len(rows)}
